"""Portable CI XLSX export when the desktop artifact-tool bundle is unavailable.

Usage: python coverage/export_ci_report.py --run artifacts/runs/<run-id>
The local reviewed workbook is authored with export_web_report.mjs.
"""
import argparse
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def export(run):
    report = json.loads((run / 'report.json').read_text(encoding='utf-8'))
    if report.get('suite') != 'web':
        raise ValueError('Expected Web suite report')
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    summary.append(['HuTube Web E2E', report['run_id']])
    summary.append(['SMOKE checks route only; REUSED does not re-execute creation.'])
    summary.append(['Unimplemented scenarios in WEB_TEST_SCENARIOS.md are still PLANNED.'])
    if report.get('run_error'):
        summary.cell(2, 2, 'INCOMPLETE RUN: ' + report['run_error'])
    summary.append(['Application', 'PASSED', 'FAILED', 'BLOCKED', 'REUSED', 'TOTAL'])
    headers = ['Test Case#', 'Test Title / Scenario', 'Test Summary', 'Test Steps', 'Test Data', 'Expected Result', 'Post-condition', 'Actual Result', 'Status', 'Notes']
    for application in ('user', 'admin'):
        name = application.title()
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for case in report['cases']:
            if case['application'] != application:
                continue
            sheet.append([case['id'], case['name'], case['module'] + ' / ' + case['actor'],
                          case.get('steps') or 'Assertions in Playwright script', case.get('data') or 'Synthetic E2E fixtures',
                          case.get('expected') or 'Assertions succeed', 'Isolated managed E2E database',
                          case.get('actual') or ('Fixture reused and verified' if case['status'] == 'reused' else 'Assertions completed' if case['status'] == 'passed' else 'See screenshot/report'),
                          case['status'].upper(), '\n'.join(filter(None, [case.get('screenshot') or 'No screenshot',
                              ('First action PASS: ' + case['first_passed_evidence']['screenshot']) if case.get('first_passed_evidence') else None,
                              ('Script: ' + case['script'] + ':' + str(case['script_line'])) if case.get('script') else None]))])
            evidence = case.get('screenshot')
            if evidence:
                sheet.cell(sheet.max_row, 10).hyperlink = '../../../' + evidence
        sheet.freeze_panes = 'B2'
        sheet.auto_filter.ref = sheet.dimensions
        table = Table(displayName=name+'WebCases', ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
        widths = [26,44,25,42,30,44,34,43,14,58]
        for index,width in enumerate(widths,1):
            sheet.column_dimensions[sheet.cell(1,index).column_letter].width=width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            sheet.row_dimensions[row[0].row].height=110 if row[0].row > 1 else 35
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='1D5278');cell.font=Font(bold=True,color='FFFFFF')
        colors={'PASSED':'E2F1E9','FAILED':'FBE4E5','BLOCKED':'FFF0D5','REUSED':'E5EEF8'}
        for row in range(2,sheet.max_row+1):
            sheet.cell(row,9).fill=PatternFill('solid',fgColor=colors[sheet.cell(row,9).value])
        end=max(2,sheet.max_row)
        line=summary.max_row+1
        summary.append([name]+[f'=COUNTIF(\'{name}\'!I2:I{end},"{status}")' for status in ('PASSED','FAILED','BLOCKED','REUSED')]+[f'=SUM(B{line}:E{line})'])
    for column in 'ABCDEF':
        summary.column_dimensions[column].width=24
    summary.freeze_panes='B5'
    destination=run/'HuTube_Web_Testcase.xlsx'
    workbook.save(destination)
    print(destination)
    return destination


if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run',type=Path)
    args=parser.parse_args()
    if args.run:
        selected=args.run
    else:
        root=Path(__file__).resolve().parents[1]
        selected=Path(json.loads((root/'artifacts/web-latest.json').read_text(encoding='utf-8'))['run_dir'])
    export(selected)
